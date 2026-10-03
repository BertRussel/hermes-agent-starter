"""Local credential-independent supporting workflows; imports are lazy.

These APIs operate only in a new caller-selected output directory. They do not
connect accounts, activate jobs, launch agents, or infer creative approval.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET
from urllib.parse import urlsplit


def fetch_public_reference(url: str, *, allowed_urls: list[str], timeout: float = 10) -> dict:
    """Retrieve explicitly allowed public text without auth, cookies or proxies.

    This is a research alternative, not a JS browser, search engine or factual
    endorsement. Plain HTTP is supported only for explicit loopback fixtures.
    Redirects require a separate approved retrieval; never follow implicitly.
    """
    from html.parser import HTMLParser
    from urllib.error import HTTPError
    from urllib.request import HTTPRedirectHandler, ProxyHandler, Request, build_opener

    parsed = urlsplit(url)
    if (parsed.username is not None or parsed.password is not None
            or any(ord(character) <= 32 for character in url)):
        raise ValueError('credential/control-bearing URL rejected')
    if (url not in allowed_urls or not parsed.hostname or parsed.fragment
            or not (parsed.scheme == 'https' or
                    (parsed.scheme == 'http' and parsed.hostname in {'localhost', '127.0.0.1', '::1'}))):
        raise ValueError('explicit public URL allowlist required')
    if not 0 < timeout <= 30:
        raise ValueError('bounded timeout required')

    class NoRedirect(HTTPRedirectHandler):
        def redirect_request(self, req, fp, code, msg, headers, newurl):
            raise ValueError('redirect requires separate approved URL')

    class Title(HTMLParser):
        def __init__(self):
            super().__init__()
            self.in_title = False
            self.parts = []

        def handle_starttag(self, tag, attrs):
            if tag == 'title':
                self.in_title = True

        def handle_endtag(self, tag):
            if tag == 'title':
                self.in_title = False

        def handle_data(self, data):
            if self.in_title:
                self.parts.append(data)

    opener = build_opener(ProxyHandler({}), NoRedirect())
    request = Request(url, headers={'Accept': 'text/html, text/plain',
                                   'User-Agent': 'public-starter-reference-canary/1'})
    with opener.open(request, timeout=timeout) as response:
        if response.status != 200:
            raise HTTPError(url, response.status, 'non-success reference response', response.headers, None)
        content_type = response.headers.get_content_type()
        if content_type not in {'text/html', 'text/plain'}:
            raise ValueError('unsupported public reference content type')
        data = response.read(1024 * 1024 + 1)
        if len(data) > 1024 * 1024:
            raise ValueError('public reference exceeds byte limit')
        text = data.decode(response.headers.get_content_charset() or 'utf-8', errors='strict')
    title = Title()
    if content_type == 'text/html':
        title.feed(text)
    return {'status': 'retrieved-not-independently-verified', 'source': url,
            'sha256': hashlib.sha256(data).hexdigest(), 'bytes': len(data),
            'title': ''.join(title.parts).strip()[:512], 'content_type': content_type,
            'browser_javascript': 'not-exercised', 'authentication': 'not-used'}


def research_feed(xml: bytes) -> list[dict[str, str]]:
    """Parse bounded local RSS input, retaining source links, without networking."""
    if len(xml) > 1024 * 1024 or b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
        raise ValueError("unsafe feed")
    root = ET.fromstring(xml)
    if root.tag != "rss":
        raise ValueError("unsupported feed format")
    entries = []
    for item in root.findall("./channel/item")[:100]:
        title, link = item.findtext("title", ""), item.findtext("link", "")
        parsed = urlsplit(link)
        if (parsed.scheme != "https" or not parsed.hostname or parsed.username is not None
                or parsed.password is not None or any(ord(character) <= 32 for character in link)):
            raise ValueError("source link must be credential-free HTTPS with a host and no controls")
        entries.append({"title": title, "source": link, "evidence": "feed declaration; not independent factual verification"})
    return entries


def render_geometric_art(design: dict, destination: Path) -> dict:
    """Render original bounded rectangle data, with no external assets or scripts."""
    import re
    from PIL import Image, ImageDraw

    if not isinstance(design, dict) or set(design) != {"width", "height", "rectangles"}:
        raise ValueError("invalid design")
    width, height = design["width"], design["height"]
    if any(type(value) is not int or not 1 <= value <= 2048 for value in (width, height)):
        raise ValueError("invalid dimensions")
    rectangles = design["rectangles"]
    if not isinstance(rectangles, list) or not 1 <= len(rectangles) <= 100:
        raise ValueError("invalid rectangles")
    for rectangle in rectangles:
        if not isinstance(rectangle, dict) or set(rectangle) != {"x", "y", "width", "height", "fill"}:
            raise ValueError("invalid rectangle")
        x, y, w, h = (rectangle[key] for key in ("x", "y", "width", "height"))
        if (any(type(value) is not int for value in (x, y, w, h)) or
                x < 0 or y < 0 or w < 1 or h < 1 or x + w > width or y + h > height):
            raise ValueError("rectangle outside canvas")
        if not isinstance(rectangle["fill"], str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", rectangle["fill"]):
            raise ValueError("invalid fill")
    if destination.exists() or destination.is_symlink():
        raise ValueError("output directory must be absent")
    destination.mkdir(mode=0o700)
    image = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    drawing = ImageDraw.Draw(image)
    svg = ET.Element("svg", {"xmlns": "http://www.w3.org/2000/svg", "width": str(width), "height": str(height)})
    for rectangle in rectangles:
        x, y, w, h = (rectangle[key] for key in ("x", "y", "width", "height"))
        drawing.rectangle((x, y, x + w - 1, y + h - 1), fill=rectangle["fill"])
        ET.SubElement(svg, "rect", {key: str(value) for key, value in rectangle.items()})
    source = destination / "editable.svg"
    source.write_bytes(ET.tostring(svg, encoding="utf-8", xml_declaration=True))
    output = destination / "render.png"
    image.save(output)
    with Image.open(output) as restored:
        assert restored.size == (width, height) and restored.getbbox() is not None
    source.chmod(0o600)
    output.chmod(0o600)
    return {"status": "rendered-not-owner-approved", "source_rights": "original-input-required",
            "sha256": {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in (source, output)}}


def media_canary(destination: Path) -> dict:
    """Generate original PCM audio, then encode/probe local video if provisioned."""
    import math
    import struct
    import wave
    import shutil
    import subprocess

    if destination.exists():
        raise ValueError("output directory must be absent")
    destination.mkdir(mode=0o700)
    audio_path = destination / "local.wav"
    with wave.open(str(audio_path), "wb") as audio:
        audio.setparams((1, 2, 8000, 0, "NONE", "not compressed"))
        audio.writeframes(b"".join(struct.pack("<h", int(1000 * math.sin(2 * math.pi * 440 * index / 8000))) for index in range(8000)))
    with wave.open(str(audio_path), "rb") as audio:
        assert audio.getnframes() == 8000 and audio.getframerate() == 8000
        assert any(audio.readframes(8000))
    audio_path.chmod(0o600)
    report = {"audio": "passed", "audio_sha256": hashlib.sha256(audio_path.read_bytes()).hexdigest(),
              "video": {"status": "blocked", "reason": "ffmpeg/ffprobe not provisioned"}}
    ffmpeg, ffprobe = shutil.which("ffmpeg"), shutil.which("ffprobe")
    if not ffmpeg or not ffprobe:
        return report
    environment = {"HOME": str(destination.absolute()), "TMPDIR": str(destination.absolute()), "PATH": "/usr/bin:/bin", "LANG": "C.UTF-8"}
    output = destination / "local.mp4"
    arguments = [ffmpeg, "-nostdin", "-n", "-loglevel", "error", "-f", "lavfi", "-i", "color=black:size=32x32:rate=1",
                 "-t", "1", "-c:v", "mpeg4", "-threads", "1", "-pix_fmt", "yuv420p", str(output)]
    encoded = subprocess.run(arguments, capture_output=True, stdin=subprocess.DEVNULL, env=environment, timeout=30)
    if encoded.returncode:
        report["video"] = {"status": "blocked", "reason": "actual ffmpeg encode failed", "exit_code": encoded.returncode}
        return report
    probe_arguments = [ffprobe, "-v", "error", "-show_entries", "stream=codec_name,width,height", "-of", "json", str(output)]
    probe = subprocess.run(probe_arguments,
                           capture_output=True, stdin=subprocess.DEVNULL, env=environment, timeout=30)
    assert probe.returncode == 0
    metadata = json.loads(probe.stdout)
    assert metadata["streams"][0]["width"] == 32 and metadata["streams"][0]["height"] == 32
    assert metadata["streams"][0]["codec_name"] == "mpeg4"
    output.chmod(0o600)
    version = subprocess.run([ffmpeg, "-version"], capture_output=True, stdin=subprocess.DEVNULL, env=environment, timeout=30)
    assert version.returncode == 0
    report["video"] = {"status": "passed", "metadata": metadata,
                       "encode_argv": arguments, "probe_argv": probe_arguments,
                       "ffmpeg_version": version.stdout.decode().splitlines()[0], "sha256": hashlib.sha256(output.read_bytes()).hexdigest()}
    return report


def local_canary(destination: Path) -> dict:
    """Create and reopen real DOCX/XLSX/PDF/raster outputs, plus local RSS."""
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from reportlab.pdfgen.canvas import Canvas
    from pypdf import PdfReader
    from PIL import Image

    if destination.exists():
        raise ValueError("output directory must be absent")
    destination.mkdir(mode=0o700)
    text = "Harmless local supporting workflow canary"
    document = Document()
    document.add_paragraph(text)
    document.save(destination / "document.docx")
    assert Document(destination / "document.docx").paragraphs[0].text == text
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active.append(["label", "value"])
    workbook.active.append(["verified", 42])
    workbook.save(destination / "workbook.xlsx")
    restored = load_workbook(destination / "workbook.xlsx")
    assert restored.active is not None
    assert restored.active["B2"].value == 42
    restored.close()
    pdf = Canvas(str(destination / "document.pdf"), invariant=1)
    pdf.drawString(36, 750, text)
    pdf.save()
    assert text in PdfReader(destination / "document.pdf").pages[0].extract_text()
    image = Image.new("RGB", (32, 32), "white")
    image.save(destination / "raster.png")
    with Image.open(destination / "raster.png") as reopened:
        assert reopened.size == (32, 32) and reopened.getpixel((0, 0)) == (255, 255, 255)
    research = research_feed(b'<rss><channel><item><title>Local citation fixture</title><link>https://example.com/source</link></item></channel></rss>')
    (destination / "research.json").write_text(json.dumps(research, indent=2) + "\n")
    files = {path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(destination.iterdir())}
    for path in destination.iterdir():
        path.chmod(0o600)
    return {"document_roundtrip": "passed", "raster_roundtrip": "passed", "local_rss": "passed", "files": files,
            "not_exercised": ["spreadsheet formula calculation", "visual editorial quality", "authenticated browser", "external research", "audio/video", "model inference"]}
